"""代码审查回归：所有网络和微信发布均使用本地替身。"""

import asyncio
import base64
import io
import ipaddress
import json
import socket
import time
import unittest
from datetime import datetime
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import httpx
from PIL import Image as PILImage
from pptx import Presentation
from langbot_plugin.api.entities.builtin.platform.message import Image, MessageChain, Plain

from components.event_listeners import diana_features as df
from components.event_listeners import whale_listener as wl


class ListenerHarness(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        state = {
            "enabled": True, "enabled_groups": ["room-A", "room-B"],
            "users": {}, "groups": {}, "voice_reply_enabled": False,
            "meme_enabled": False, "speak_up_enabled": True,
        }
        with patch.object(wl.WhaleListener, "_load_persistent_state", return_value=state), \
                patch.object(wl.WhaleListener, "_load_bundled_memes", return_value={}), \
                patch.object(df, "load_stolen_memes", return_value=[]):
            self.listener = wl.WhaleListener()
        obj = self.listener
        obj._save_persistent_state = lambda: None
        obj._get_config_str = lambda key, default: {
            "deepseek_api_key": "test-key", "deepseek_base_url": "https://model.invalid",
            "target_bot_uuid": "",
        }.get(key, default)
        obj._resolve_bot_uuid = AsyncMock(return_value="bot-A")
        obj._get_balance_str = AsyncMock(return_value="10")
        obj._async_post_group_turn_judge = AsyncMock()
        obj._async_post_turn_judge = AsyncMock()
        obj._split_reply_bubbles = lambda text, **kwargs: [text]
        self.requests = []
        self.model_reply = "测试回复"
        original_client = httpx.AsyncClient

        def respond(req):
            self.requests.append(req)
            if req.url.path.endswith("/chat/completions"):
                return httpx.Response(200, json={
                    "choices": [{"message": {"content": self.model_reply}}],
                    "usage": {"total_tokens": 1},
                })
            raise AssertionError(f"未预期的外部请求：{req.url}")

        self._patch(httpx, "AsyncClient", side_effect=lambda **kwargs: original_client(
            transport=httpx.MockTransport(respond), **kwargs))
        self._patch(df, "resolve_user_display_name", AsyncMock(return_value="测试用户"))
        self._patch(df, "refresh_dynamic_bot_names", AsyncMock())
        self.publisher = self._patch(df, "publish_wcf_moment", AsyncMock(return_value=True))

    def _patch(self, target, name, *args, **kwargs):
        item = patch.object(target, name, *args, **kwargs)
        self.addCleanup(item.stop)
        return item.start()

    async def ask(self, text, *, group=False, room="wxid_test", sender="wxid_test", image=None):
        comps = [Plain(text=text)] if text else []
        if image:
            comps.append(Image(base64=image))
        ctx = SimpleNamespace(
            event=SimpleNamespace(
                message_chain=MessageChain(comps), sender_id=sender, launcher_id=room,
                source_platform_object={}, reply_message_chain=None,
            ),
            query_id=None, is_prevented_default=lambda: False, prevent_default=lambda: None,
        )
        self.requests.clear()
        await self.listener._handle_chat(ctx, is_group=group)
        await asyncio.sleep(0)
        return ctx

    async def upload(self, secret, **kwargs):
        data = base64.b64encode(secret.encode()).decode()
        await self.ask(
            f"<msg><appmsg><title>notes.txt</title><type>6</type>"
            f"<file_b64>{data}</file_b64></appmsg></msg>", **kwargs,
        )

    def payload(self):
        req = next(r for r in self.requests if r.url.path.endswith("/chat/completions"))
        return json.loads(req.content)

    def prompt(self):
        return json.dumps(self.payload()["messages"], ensure_ascii=False)


class ListenerRegressions(ListenerHarness):
    async def test_private_document_is_not_recalled_in_group(self):
        await self.upload("PRIVATE-DOC-SENTINEL")
        await self.ask("大肥鱼 总结刚才的文件", group=True, room="room-A")
        self.assertNotIn("PRIVATE-DOC-SENTINEL", self.prompt())

    async def test_group_document_is_not_recalled_in_another_group(self):
        await self.upload("GROUP-A-SENTINEL", group=True, room="room-A")
        await self.ask("大肥鱼 总结刚才的文件", group=True, room="room-B")
        self.assertNotIn("GROUP-A-SENTINEL", self.prompt())

    async def test_group_document_is_not_recalled_in_private(self):
        await self.upload("GROUP-DOC-SENTINEL", group=True, room="room-A")
        await self.ask("总结刚才的文件")
        self.assertNotIn("GROUP-DOC-SENTINEL", self.prompt())

    async def test_document_is_not_recalled_by_another_bot(self):
        await self.upload("BOT-A-SENTINEL")
        self.listener._resolve_bot_uuid.return_value = "bot-B"
        await self.ask("总结刚才的文件")
        self.assertNotIn("BOT-A-SENTINEL", self.prompt())

    async def test_same_session_document_followup_still_works(self):
        await self.upload("SAME-SESSION-SENTINEL", group=True, room="room-A")
        await self.ask("大肥鱼 总结刚才的文件", group=True, room="room-A")
        self.assertIn("SAME-SESSION-SENTINEL", self.prompt())

    async def test_document_is_not_recalled_by_another_member(self):
        await self.upload("MEMBER-A-SENTINEL", group=True, room="room-A")
        await self.ask("大肥鱼 总结刚才的文件", group=True, room="room-A", sender="wxid_other")
        self.assertNotIn("MEMBER-A-SENTINEL", self.prompt())

    async def test_reset_clears_only_current_session_document(self):
        await self.upload("ROOM-A-SENTINEL", group=True, room="room-A")
        await self.upload("ROOM-B-SENTINEL", group=True, room="room-B")
        await self.ask("大肥鱼 重置", group=True, room="room-A")
        await self.ask("大肥鱼 总结刚才的文件", group=True, room="room-A")
        self.assertNotIn("ROOM-A-SENTINEL", self.prompt())
        await self.ask("大肥鱼 总结刚才的文件", group=True, room="room-B")
        self.assertIn("ROOM-B-SENTINEL", self.prompt())

    async def test_expired_document_is_not_recalled(self):
        await self.upload("EXPIRED-SENTINEL")
        for doc in self.listener._user_docs.values():
            doc["ts"] = time.time() - 1801
        await self.ask("总结刚才的文件")
        self.assertNotIn("EXPIRED-SENTINEL", self.prompt())

    async def test_group_image_with_text_reaches_model(self):
        buf = io.BytesIO()
        PILImage.new("RGB", (2, 2)).save(buf, format="JPEG")
        await self.ask("大肥鱼 看看这张图", group=True, room="room-A",
                       image=base64.b64encode(buf.getvalue()).decode())
        content = self.payload()["messages"][-1]["content"]
        self.assertIsInstance(content, list)
        self.assertTrue(any(item["type"] == "image_url" for item in content))
        self.assertIn("room-A", self.listener._group_recent_images)

    async def test_message_image_url_does_not_request_loopback(self):
        chain = MessageChain([Image(url="http://127.0.0.1/private.jpg")])
        result = await self.listener._extract_image_b64_from_chain(chain)
        self.assertIsNone(result)
        self.assertEqual(self.requests, [])

    async def test_sticker_vision_url_does_not_request_loopback(self):
        result = await self.listener._fetch_sticker_strip_b64('<emoji cdnurl="http://127.0.0.1/a.gif" />')
        self.assertIsNone(result)
        self.assertEqual(self.requests, [])

    async def test_public_moment_does_not_use_private_memory_or_mood_reason(self):
        self.listener._state["users"]["wxid_test"] = {
            "hasPrivateChat": True, "facts": ["PRIVATE-FACT-SENTINEL"],
            "archive": {"一": ["PRIVATE-ARCHIVE-SENTINEL"]},
            "events": [{"text": "PRIVATE-EVENT-SENTINEL"}],
            "lastReason": "PRIVATE-REASON-SENTINEL",
        }
        await self.ask("大肥鱼 发朋友圈 今天天气不错")
        prompt = self.prompt()
        for marker in ("PRIVATE-FACT", "PRIVATE-ARCHIVE", "PRIVATE-EVENT", "PRIVATE-REASON"):
            self.assertNotIn(marker, prompt)
        self.assertIn("今天天气不错", prompt)
        self.publisher.assert_awaited_once()

    async def one_loop_iteration(self, loop):
        fixed = SimpleNamespace(now=lambda _: datetime(2026, 10, 27, 22, 0, tzinfo=wl.TZ))
        with patch.object(wl.asyncio, "sleep", AsyncMock(side_effect=[None, asyncio.CancelledError()])), \
                patch.object(wl, "datetime", fixed), \
                patch.object(wl.random, "uniform", return_value=0.0), \
                patch.object(wl.random, "random", return_value=0.0):
            with self.assertRaises(asyncio.CancelledError):
                await loop()

    async def test_speak_up_does_not_publish_moments_when_disabled(self):
        obj = self.listener
        obj._state["auto_moments_enabled"] = False
        obj._state["users"]["wxid_test"] = {
            "hasPrivateChat": True, "lastChatAt": int(time.time() * 1000) - 3600000,
        }
        obj._generate_and_post_moment = AsyncMock(return_value=("测试动态", None))
        obj._send_proactive_wechat = AsyncMock()
        obj._trigger_speak_up_for_user = AsyncMock()
        await self.one_loop_iteration(obj._speak_up_loop)
        obj._generate_and_post_moment.assert_not_awaited()
        obj._trigger_speak_up_for_user.assert_awaited_once_with("wxid_test", manual=False)

    async def test_daily_moments_does_not_select_a_private_user(self):
        obj = self.listener
        obj._state.update({
            "auto_moments_enabled": True, "autoMomentTargetDate": "2026-10-27",
            "autoMomentTargetMin": 1290,
            "users": {"wxid_test": {"hasPrivateChat": True}},
        })
        obj._generate_and_post_moment = AsyncMock(return_value=("测试动态", None))
        await self.one_loop_iteration(obj._auto_moments_loop)
        obj._generate_and_post_moment.assert_awaited_once_with()

    async def test_private_mood_judge_does_not_publish_private_reason(self):
        obj = self.listener
        obj._state["users"]["wxid_test"] = {"hasPrivateChat": True}
        obj._generate_and_post_moment = AsyncMock(return_value=("测试动态", None))
        obj._send_proactive_wechat = AsyncMock()
        self.model_reply = json.dumps({
            "moodDelta": 4, "confidence": 90, "reason": "PRIVATE-MOOD-SENTINEL",
            "memory_fact": "PRIVATE-FACT-SENTINEL",
        })
        await wl.WhaleListener._async_post_turn_judge(
            obj, api_key="test-key", base_url="https://model.invalid", model="deepseek-chat",
            user_id="wxid_test", user_text="测试私聊", reply_text="测试回复",
        )
        self.assertEqual(obj._state["users"]["wxid_test"]["lastReason"], "PRIVATE-MOOD-SENTINEL")
        obj._generate_and_post_moment.assert_not_awaited()
        obj._send_proactive_wechat.assert_not_awaited()

    async def test_daily_moments_honors_disabled_flag(self):
        self.listener._state["auto_moments_enabled"] = False
        self.listener._generate_and_post_moment = AsyncMock()
        await self.one_loop_iteration(self.listener._auto_moments_loop)
        self.listener._generate_and_post_moment.assert_not_awaited()

    async def test_daily_moments_does_not_repeat_on_same_day(self):
        self.listener._state["lastAutoMomentDate"] = "2026-10-27"
        self.listener._generate_and_post_moment = AsyncMock()
        await self.one_loop_iteration(self.listener._auto_moments_loop)
        self.listener._generate_and_post_moment.assert_not_awaited()

    async def test_private_sticker_placeholder_is_not_ignored(self):
        ctx = await self.ask("[动画表情]")
        self.assertIsNotNone(ctx.event.reply_message_chain)
        self.assertIn("微信表情包", self.prompt())

    async def test_private_voice_placeholder_is_not_ignored(self):
        ctx = await self.ask("[语音]")
        self.assertIsNotNone(ctx.event.reply_message_chain)
        self.assertIn("微信语音", self.prompt())

    async def test_private_voice_xml_with_audio_b64_is_transcribed_and_replied(self):
        with patch.object(
            self.listener,
            "_transcribe_incoming_voice",
            AsyncMock(return_value=("今天中午吃什么呀", "开心", 3.1)),
        ):
            ctx = await self.ask('<msg><voicemsg audio_b64="UklGRg==" voicelength="3120" length="5529" /></msg>')
        self.assertIsNotNone(ctx.event.reply_message_chain)
        self.assertIn("今天中午吃什么呀", self.prompt())
        self.assertIn("开心", self.prompt())


class PublicUrlRegressions(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.requests = []
        self.response_for = lambda req: httpx.Response(200, text="PUBLIC-CONTENT")
        original_client = httpx.AsyncClient

        def respond(req):
            self.requests.append(req)
            if req.url.host == "r.jina.ai":
                return httpx.Response(503)
            return self.response_for(req)

        def resolve(host, *args, **kwargs):
            try:
                address = str(ipaddress.ip_address(host))
            except ValueError:
                address = "8.8.8.8"
            return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", (address, 0))]

        client_patch = patch.object(httpx, "AsyncClient", side_effect=lambda **kwargs: original_client(
            transport=httpx.MockTransport(respond), **kwargs))
        dns_patch = patch.object(df.socket, "getaddrinfo", side_effect=resolve)
        client_patch.start()
        dns_patch.start()
        self.addCleanup(client_patch.stop)
        self.addCleanup(dns_patch.stop)

    async def test_platform_name_in_path_does_not_trigger_parsing(self):
        url = "http://127.0.0.1:18080/xiaohongshu.com/"
        card, image = await df.extract_and_parse_link(url)
        self.assertIsNone(card)
        self.assertIsNone(image)
        self.assertEqual(self.requests, [])
        self.assertFalse(df.has_supported_media_link_or_card(url))

    async def test_platform_dispatch_uses_real_hostname(self):
        cases = [
            ("https://www.bilibili.com/video/BV1234567890", "_parse_bilibili_url"),
            ("https://b23.tv/short", "_parse_bilibili_url"),
            ("https://x.com/person", "_parse_twitter_url"),
            ("https://m.youtube.com/watch?v=123", "_parse_youtube_or_social_url"),
            ("https://www.xiaohongshu.com/explore/123", "_parse_youtube_or_social_url"),
        ]
        for url, parser in cases:
            with self.subTest(url=url), patch.object(df, parser, AsyncMock(return_value=("CARD", None))) as call:
                self.assertTrue(df.has_supported_media_link_or_card(url))
                self.assertEqual(await df.extract_and_parse_link(url), ("CARD", None))
                call.assert_awaited_once_with(url)

    async def test_hostname_spoofs_are_not_supported(self):
        for url in ("https://x.com.attacker.invalid/", "https://attacker.invalid/?next=x.com/",
                    "https://xiaohongshu.com@127.0.0.1/", "https://notxiaohongshu.com/"):
            with self.subTest(url=url):
                self.assertFalse(df.has_supported_media_link_or_card(url))

    async def test_social_redirect_cannot_reach_loopback(self):
        self.response_for = lambda req: httpx.Response(302, headers={"Location": "http://127.0.0.1/private"})
        await df.extract_and_parse_link("https://www.xiaohongshu.com/explore/1")
        self.assertTrue(self.requests)
        self.assertNotIn("127.0.0.1", [req.url.host for req in self.requests])

    async def test_short_link_redirect_cannot_reach_private_ip(self):
        self.response_for = lambda req: httpx.Response(302, headers={"Location": "http://192.168.1.1/private"})
        await df.extract_and_parse_link("https://b23.tv/short")
        self.assertNotIn("192.168.1.1", [req.url.host for req in self.requests])

    async def test_webpage_redirect_cannot_reach_cloud_metadata(self):
        self.response_for = lambda req: httpx.Response(302, headers={"Location": "http://169.254.169.254/latest"})
        await df.read_webpage_content("https://public.invalid/start")
        self.assertNotIn("169.254.169.254", [req.url.host for req in self.requests])

    async def test_webpage_public_relative_redirect_still_works(self):
        self.response_for = lambda req: (
            httpx.Response(302, headers={"Location": "/done"}) if req.url.path == "/start"
            else httpx.Response(200, text="PUBLIC-REDIRECT-SENTINEL")
        )
        result = await df.read_webpage_content("https://public.invalid/start")
        self.assertIn("PUBLIC-REDIRECT-SENTINEL", result)

    async def test_cover_download_rejects_private_addresses(self):
        for url in ("http://127.0.0.1/a.jpg", "http://10.0.0.1/a.jpg", "http://[::1]/a.jpg"):
            with self.subTest(url=url):
                self.assertIsNone(await df._download_image_as_jpeg_b64(url))
        self.assertEqual(self.requests, [])

    async def test_cover_redirect_does_not_reach_private_address(self):
        self.response_for = lambda req: httpx.Response(302, headers={"Location": "http://10.0.0.1/a.jpg"})
        self.assertIsNone(await df._download_image_as_jpeg_b64("https://cdn.invalid/a.jpg"))
        self.assertNotIn("10.0.0.1", [req.url.host for req in self.requests])

    async def test_stolen_sticker_does_not_request_private_address(self):
        result = await df.steal_incoming_sticker(
            '<emoji cdnurl="http://10.0.0.1/a.gif" />', {}, [],
        )
        self.assertFalse(result)
        self.assertEqual(self.requests, [])

    async def test_public_cover_download_still_works(self):
        buf = io.BytesIO()
        PILImage.new("RGB", (2, 2)).save(buf, format="PNG")
        self.response_for = lambda req: httpx.Response(200, content=buf.getvalue())
        result = await df._download_image_as_jpeg_b64("https://cdn.invalid/a.png")
        self.assertIsNotNone(result)
        self.assertEqual(PILImage.open(io.BytesIO(base64.b64decode(result))).format, "JPEG")

    async def test_dns_private_address_is_blocked(self):
        answer = [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("10.0.0.1", 0))]
        with patch.object(df.socket, "getaddrinfo", return_value=answer):
            await df.extract_and_parse_link("https://www.xiaohongshu.com/explore/1")
        self.assertEqual(self.requests, [])

    def test_empty_dns_answer_is_not_public(self):
        with patch.object(df.socket, "getaddrinfo", return_value=[]):
            self.assertFalse(df._is_safe_public_url("https://public.invalid/"))


class PowerPointRegressions(unittest.TestCase):
    def parse(self, slide_count):
        prs = Presentation()
        for index in range(slide_count):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = f"SLIDE-{index:02d}-SENTINEL"
        buf = io.BytesIO()
        prs.save(buf)
        xml = "<file_name>demo.pptx</file_name><file_b64>" + base64.b64encode(buf.getvalue()).decode() + "</file_b64>"
        return df.parse_wechat_file_xml(xml)

    def test_valid_pptx_extracts_slide_text(self):
        name, text = self.parse(1)
        self.assertEqual(name, "demo.pptx")
        self.assertIn("SLIDE-00-SENTINEL", text)

    def test_pptx_stops_after_forty_slides(self):
        _, text = self.parse(41)
        self.assertIn("SLIDE-39-SENTINEL", text)
        self.assertNotIn("SLIDE-40-SENTINEL", text)


if __name__ == "__main__":
    unittest.main()
