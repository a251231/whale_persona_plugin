"""Diana 审查回归：限制工作表迭代，并按机器人及完整群 ID 隔离日志。"""

import base64
import inspect
import io
import json
import unittest
from unittest.mock import patch

import httpx
import openpyxl
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from components.event_listeners import diana_features as df


def workbook_message(*cells):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "测试工作表"
    for row, column, value in cells:
        sheet.cell(row, column, value)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    encoded = base64.b64encode(output.getvalue()).decode("ascii")
    return f"<file_name>表格.xlsx</file_name><file_b64>{encoded}</file_b64>"


class WorkbookBoundsTests(unittest.TestCase):
    def parse_with_iteration_counts(self, message):
        counts = {"rows": 0, "max_columns": 0}
        original = ReadOnlyWorksheet.iter_rows

        def counted(sheet, *args, **kwargs):
            for row in original(sheet, *args, **kwargs):
                counts["rows"] += 1
                counts["max_columns"] = max(counts["max_columns"], len(row))
                yield row

        with patch.object(ReadOnlyWorksheet, "iter_rows", counted):
            result = df.parse_wechat_file_xml(message)
        return result, counts

    def test_sparse_workbook_iterates_at_most_first_80_rows(self):
        message = workbook_message((1, 1, "首行内容"), (80, 1, "第80行"), (1000, 1, "末行内容"))
        (_, text), counts = self.parse_with_iteration_counts(message)
        self.assertLessEqual(counts["rows"], 80)
        self.assertIn("首行内容", text)
        self.assertIn("第80行", text)
        self.assertNotIn("末行内容", text)

    def test_sparse_workbook_iterates_at_most_first_80_columns(self):
        message = workbook_message((1, 1, "首列内容"), (1, 80, "第80列"), (1, 160, "末列内容"))
        (_, text), counts = self.parse_with_iteration_counts(message)
        self.assertLessEqual(counts["max_columns"], 80)
        self.assertIn("首列内容", text)
        self.assertIn("第80列", text)
        self.assertNotIn("末列内容", text)

    def test_normal_workbook_keeps_sheet_name_and_cell_values(self):
        filename, text = df.parse_wechat_file_xml(workbook_message((1, 1, "姓名"), (1, 2, "数量"), (2, 1, "鱼"), (2, 2, 3)))
        self.assertEqual(filename, "表格.xlsx")
        self.assertIn("【工作表：测试工作表】", text)
        self.assertIn("姓名 | 数量", text)
        self.assertIn("鱼 | 3", text)


class InterjectLogTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.state = {}
        self.reason = "当前群测试原因"
        original_client = httpx.AsyncClient

        def respond(request):
            self.assertEqual(request.url.host, "model.invalid")
            return httpx.Response(200, json={
                "choices": [{"message": {"content": json.dumps({
                    "should_reply": True, "reply": "测试回复", "reason": self.reason,
                }, ensure_ascii=False)}}],
                "usage": {"total_tokens": 3},
            })

        client_patch = patch.object(httpx, "AsyncClient", side_effect=lambda **kwargs: original_client(
            transport=httpx.MockTransport(respond), **kwargs))
        client_patch.start()
        self.addCleanup(client_patch.stop)

    def format_logs(self, *, bot_uuid, group_id):
        self.assertIn("bot_uuid", inspect.signature(df.format_interject_logs).parameters)
        self.assertIn("group_id", inspect.signature(df.format_interject_logs).parameters)
        return df.format_interject_logs(self.state, bot_uuid=bot_uuid, group_id=group_id)

    async def produce_log(self, producer, *, bot_uuid="bot-A", group_id="1234567890-room-A"):
        self.assertIn("bot_uuid", inspect.signature(producer).parameters)
        kwargs = dict(api_key="test-key", base_url="https://model.invalid", model="test-model",
                      bot_uuid=bot_uuid, group_id=group_id, recent_lines=["群友：聊聊测试"], state=self.state)
        if producer is df.evaluate_group_active_continuation:
            kwargs.update(sender_id="sender-A", sender_name="群友", current_text="继续聊测试",
                          last_bot_reply="刚刚的回复", is_same_speaker=True)
        result = await producer(**kwargs)
        self.assertTrue(result[0])
        return self.state["interject_logs"][-1]

    async def test_interject_produces_readable_log_with_full_scope(self):
        entry = await self.produce_log(df.evaluate_group_interject)
        self.assertEqual(entry["bot_uuid"], "bot-A")
        self.assertEqual(entry["group_id"], "1234567890-room-A")
        self.assertIn(self.reason, self.format_logs(bot_uuid="bot-A", group_id="1234567890-room-A"))

    async def test_continuation_produces_readable_log_with_full_scope(self):
        entry = await self.produce_log(df.evaluate_group_active_continuation)
        self.assertEqual(entry["bot_uuid"], "bot-A")
        self.assertEqual(entry["group_id"], "1234567890-room-A")
        self.assertIn(self.reason, self.format_logs(bot_uuid="bot-A", group_id="1234567890-room-A"))

    async def test_group_ids_with_same_prefix_do_not_share_logs(self):
        await self.produce_log(df.evaluate_group_interject)
        self.assertNotIn(self.reason, self.format_logs(bot_uuid="bot-A", group_id="1234567890-room-B"))

    async def test_another_robot_cannot_read_group_log(self):
        await self.produce_log(df.evaluate_group_active_continuation)
        self.assertNotIn(self.reason, self.format_logs(bot_uuid="bot-B", group_id="1234567890-room-A"))

    async def test_all_groups_query_only_returns_current_robot(self):
        for bot, group, reason in [("bot-A", "room-A", "机器人A群一"), ("bot-A", "room-B", "机器人A群二"),
                                   ("bot-B", "room-A", "机器人B隐私")]:
            self.reason = reason
            await self.produce_log(df.evaluate_group_interject, bot_uuid=bot, group_id=group)
        text = self.format_logs(bot_uuid="bot-A", group_id=None)
        self.assertIn("机器人A群一", text)
        self.assertIn("机器人A群二", text)
        self.assertNotIn("机器人B隐私", text)

    def test_legacy_and_unknown_scope_logs_are_hidden_for_every_query(self):
        self.state["interject_logs"] = [
            {"group": "1234567890", "reason": "旧版未知日志"},
            {"bot_uuid": "bot-A", "group": "1234567890", "reason": "缺完整群日志"},
            {"group_id": "1234567890-room-A", "reason": "缺机器人日志"},
            {"bot_uuid": "", "group_id": "room-A", "reason": "空机器人日志"},
            {"bot_uuid": "bot-A", "group_id": "", "reason": "空群日志"},
        ]
        for bot in ("bot-A", ""):
            for group in ("1234567890-room-A", None):
                with self.subTest(bot=bot, group=group):
                    self.assertIn("暂无群聊插话评估记录", self.format_logs(bot_uuid=bot, group_id=group))

    async def test_filtering_happens_before_recent_eight_limit(self):
        await self.produce_log(df.evaluate_group_interject)
        self.state["interject_logs"].extend([
            {"bot_uuid": "bot-B", "group_id": "room-B", "reason": "其他机器人日志"}
            for _ in range(10)
        ])
        self.assertIn(self.reason, self.format_logs(bot_uuid="bot-A", group_id="1234567890-room-A"))


if __name__ == "__main__":
    unittest.main()
